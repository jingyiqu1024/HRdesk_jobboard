from openpyxl import load_workbook
from .common import XLSX_PATH

def read_sheet(wb, name):
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else '' for h in next(rows)]
    out = []
    for row in rows:
        if not row or all(v is None for v in row): continue
        if row[0] == 'id': continue
        out.append(dict(zip(header, row)))
    return out

def load_raw():
    wb = load_workbook(str(XLSX_PATH), read_only=True)
    raw = {k: read_sheet(wb, v) for k, v in {
        'bytedance':'jd_bytedance_jobs','taotian':'jd_alibaba_taotian_jobs','campus':'jd_alibaba_campus_jobs',
        'tencent':'jd_tencent_jobs','meituan':'jd_meituan_jobs','xhs':'jd_xiaohongshu_jobs','jd':'jd_jd_jobs','pdd':'jd_pinduoduo_jobs','ai':'jd_ai_jobs'
    }.items()}
    wb.close()
    raw['taotian_all'] = raw['taotian'] + [dict(r, recruit_type='campus') for r in raw['campus']]
    return raw
